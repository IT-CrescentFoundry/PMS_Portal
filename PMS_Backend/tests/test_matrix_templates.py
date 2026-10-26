"""Generated Excel templates: metadata-driven columns, round trip, versions, safety."""
from dataclasses import replace
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models import AuditEvent, Matrix, MatrixVersion, MatrixVersionSource
from app.models.enums import KpiDirection, ScopeType
from app.services import matrix_rules, matrix_service, matrix_types
from app.services.matrix_types import DataType, FieldSpec, Source

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GLOBAL, PLANT, DEPARTMENT = ScopeType.GLOBAL, ScopeType.PLANT, ScopeType.DEPARTMENT
BANDS = matrix_rules.definition_from_bands(
    KpiDirection.HIGHER_IS_BETTER, [{"threshold": 100, "score": 5}, {"threshold": 90, "score": 4}],
    1, None, None)


@pytest.fixture
def people(as_role, org):
    return {
        "hr_head": as_role("tst.tp.hrh", [("HR_HEAD", [(GLOBAL, None, None)])]),
        "hr_head_p1": as_role("tst.tp.hrh1", [("HR_HEAD", [(PLANT, org.plants["P1"], None)])]),
        "chro": as_role("tst.tp.chro", [("CHRO", [(GLOBAL, None, None)])]),
        "dept_head": as_role("tst.tp.dh", [("DEPT_HEAD", [(DEPARTMENT, None, org.departments["Q"])])]),
    }


def load(data: bytes):
    return load_workbook(BytesIO(data))


def save(wb) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def keys(ws) -> dict[str, int]:
    return {cell.value: cell.column for cell in ws[1] if cell.value}


def put(ws, row: int, values: dict[str, object]) -> None:
    columns = keys(ws)
    for key, value in values.items():
        ws.cell(row=row, column=columns[key], value=value)


def setting(ws, key: str, value: object) -> None:
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == key)
    ws.cell(row=row, column=3).value = value  # cell(..., value=None) would leave it unchanged


def send(client, step: str, data: bytes, name: str = "matrix.xlsx"):
    return client.post(f"/api/matrices/uploads/{step}", files={"file": (name, data, XLSX)})


def type_template(client, code="KPI_SCORING") -> bytes:
    resp = client.get(f"/api/matrix-types/{code}/template")
    assert resp.status_code == 200 and resp.headers["content-type"] == XLSX
    return resp.content


def filled_type_template(client, name="Uploaded matrix") -> bytes:
    wb = load(type_template(client))
    settings = wb["Matrix"]
    setting(settings, "name", name)
    setting(settings, "attr:direction", "HIGHER_IS_BETTER")
    setting(settings, "effective_from", "2026-04-01")
    setting(settings, "no_match", "Use the default output")
    setting(settings, "default:score", 1)
    rules_sheet = wb["Rules"]
    put(rules_sheet, 3, {"rule:priority": 1, "rule:label": "Above target",
                         "cond:achievement_percent:min": 100,
                         "cond:achievement_percent:min_inclusive": "more than (>)", "out:score": 5})
    put(rules_sheet, 4, {"rule:priority": 2, "cond:achievement_percent:min": 90,
                         "cond:achievement_percent:min_inclusive": "more than (>)",
                         "cond:achievement_percent:max": 100,
                         "cond:achievement_percent:max_inclusive": "at most (≤)", "out:score": 4})
    put(wb["Test cases"], 3, {"case:label": "On target", "in:achievement_percent": 100,
                              "expect:score": 4})
    return save(wb)


def counts(db):
    return (db.scalar(select(func.count()).select_from(Matrix)),
            db.scalar(select(func.count()).select_from(MatrixVersion)))


def test_type_template_is_generated_from_metadata(people):
    wb = load(type_template(people["hr_head"][1]))
    assert wb.sheetnames == ["Instructions", "Matrix", "Applies to", "Fields", "Rules", "Test cases",
                             "Codes", "_pms"]
    assert wb["_pms"].sheet_state == "veryHidden"
    meta = {r[0].value: r[1].value for r in wb["_pms"].iter_rows()}
    assert meta["template_kind"] == "matrix_configuration" and meta["matrix_type"] == "KPI_SCORING"
    assert meta["schema_version"] == 1 and meta["template_version"] == 1
    fields = [row for row in wb["Fields"].iter_rows(min_row=3, values_only=True)]
    standard = [f.key for f in matrix_types.KPI_SCORING.standard_inputs]
    assert [r[1] for r in fields] == [*standard, "score"]
    assert fields[0][0] == "Yes" and fields[1][0] == "No"   # first standard input in use
    rule_keys = keys(wb["Rules"])
    for key in standard:
        assert f"cond:{key}:min" in rule_keys and f"cond:{key}:max_inclusive" in rule_keys
    assert "out:score" in rule_keys and wb["Rules"].row_dimensions[1].hidden
    assert wb["Rules"].data_validations.dataValidation  # bound words, Yes/No lists
    assert people["dept_head"][1].get("/api/matrix-types/KPI_SCORING/template").status_code == 403


def test_a_new_standard_field_reaches_template_and_upload(people, monkeypatch, db):
    extended = replace(matrix_types.KPI_SCORING, standard_inputs=(
        *matrix_types.KPI_SCORING.standard_inputs,
        FieldSpec("defect_ppm", "Defects per million", DataType.DECIMAL, "input", unit="ppm",
                  min="0", source=Source.ACTUAL)))
    monkeypatch.setitem(matrix_types.REGISTRY, "KPI_SCORING", extended)
    client = people["hr_head"][1]
    wb = load(type_template(client))
    assert "cond:defect_ppm:min" in keys(wb["Rules"]) and "in:defect_ppm" in keys(wb["Test cases"])
    fields = wb["Fields"]
    for row in range(3, fields.max_row + 1):  # use the new field instead of achievement %
        key = fields.cell(row=row, column=2).value
        if key in ("achievement_percent", "defect_ppm"):
            fields.cell(row=row, column=1, value="Yes" if key == "defect_ppm" else "No")
    setting(wb["Matrix"], "name", "PPM matrix")
    setting(wb["Matrix"], "attr:direction", "LOWER_IS_BETTER")
    setting(wb["Matrix"], "effective_from", "2026-04-01")
    setting(wb["Matrix"], "no_match", "Use the default output")
    setting(wb["Matrix"], "default:score", 1)
    put(wb["Rules"], 3, {"rule:priority": 1, "cond:defect_ppm:max": 100,
                         "cond:defect_ppm:max_inclusive": "at most (≤)", "out:score": 5})
    preview = send(client, "preview", save(wb)).json()
    assert preview["problems"] == [] and preview["ready"]
    assert [f["key"] for f in preview["fields"]] == ["defect_ppm", "score"]
    assert preview["rules"][0]["description"] == "Defects per million is at most 100 ppm"


def test_round_trip_of_a_version_template(people, db):
    client, chro = people["hr_head"][1], people["chro"][1]
    matrix = client.post("/api/matrices", json={
        "matrix_type": "KPI_SCORING", "name": "Round trip", "attributes": {"direction": "HIGHER_IS_BETTER"},
        "applicability": [{}]}).json()
    vid = matrix["latest_version"]["id"]
    client.put(f"/api/matrices/{matrix['id']}/versions/{vid}", json={
        "definition": BANDS, "effective_from": "2026-04-01", "applicability": [{}]})
    for step, who in (("submit", client), ("approve", chro), ("publish", client)):
        assert who.post(f"/api/matrices/{matrix['id']}/versions/{vid}/{step}", json={}).status_code == 200
    data = client.get(f"/api/matrices/{matrix['id']}/versions/{vid}/template").content
    before = counts(db)
    preview = send(client, "preview", data, "round-trip.xlsx").json()
    assert counts(db) == before                     # validate-only: nothing written
    assert preview["action"] == "new_version" and preview["problems"] == [] and preview["ready"]
    assert [r["description"] for r in preview["rules"]] == [
        "Achievement % is more than 100%", "Achievement % is more than 90% and at most 100%"]
    confirmed = send(client, "confirm", data, "round-trip.xlsx")
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["version_no"] == 2
    version = db.get(MatrixVersion, confirmed.json()["version_id"])
    assert version.source == MatrixVersionSource.UPLOAD and version.template_version == 1
    assert version.source_file_name == "round-trip.xlsx" and len(version.source_file_sha256) == 64
    original = matrix_rules.frozen_definition(BANDS)
    uploaded = matrix_rules.frozen_definition(version.definition)
    for value in ("120", "100", "95", "90", "10"):
        assert (matrix_rules.evaluate(uploaded, {"achievement_percent": value}).outputs
                == matrix_rules.evaluate(original, {"achievement_percent": value}).outputs)
    event = db.scalars(select(AuditEvent).where(AuditEvent.action == "matrix_version.upload")).one()
    assert event.changes["upload"]["file"] == "round-trip.xlsx"
    assert event.changes["upload"]["template_version"] == 1


def test_a_filled_type_template_creates_a_matrix(people, db):
    client = people["hr_head"][1]
    data = filled_type_template(client)
    preview = send(client, "preview", data).json()
    assert preview["action"] == "create" and preview["ready"], preview["problems"]
    assert preview["test_results"][0]["passed"]
    confirmed = send(client, "confirm", data).json()
    detail = client.get(f"/api/matrices/{confirmed['matrix_id']}").json()
    assert detail["name"] == "Uploaded matrix" and detail["latest_version"]["status"] == "DRAFT"
    assert detail["policy"]["code"].startswith("POL-")


@pytest.mark.parametrize("meta_key, value, expected", [
    ("template_version", 0, "template layout 0, which is no longer supported"),
    ("schema_version", 0, "incompatible version"),
    ("matrix_type", "PAYROLL", "Unknown matrix type"),
])
def test_incompatible_workbooks_are_rejected(people, meta_key, value, expected):
    client = people["hr_head"][1]
    wb = load(filled_type_template(client))
    meta = wb["_pms"]
    row = next(r for r in range(1, meta.max_row + 1) if meta.cell(row=r, column=1).value == meta_key)
    meta.cell(row=row, column=2, value=value)
    preview = send(client, "preview", save(wb)).json()
    assert not preview["ready"] and any(expected in p["message"] for p in preview["problems"])


def test_an_older_schema_is_converted_only_by_an_explicit_migration(people, monkeypatch):
    client = people["hr_head"][1]
    data = filled_type_template(client)
    newer = replace(matrix_types.KPI_SCORING, schema_version=2, migrations={1: lambda d: d})
    monkeypatch.setitem(matrix_types.REGISTRY, "KPI_SCORING", newer)
    preview = send(client, "preview", data).json()
    assert preview["ready"] and any("Converted from schema version 1 to 2" in w["message"]
                                    for w in preview["warnings"])


def test_unreadable_or_foreign_files(people):
    client = people["hr_head"][1]
    wb = load(filled_type_template(client))
    del wb["_pms"]
    resp = send(client, "preview", save(wb))
    assert resp.status_code == 422 and "not a PMS matrix template" in resp.json()["issues"][0]["message"]
    assert send(client, "preview", b"PK\x03\x04 broken").status_code == 422
    assert send(client, "preview", b"a,b\n1,2", "matrix.csv").status_code == 422
    assert send(people["dept_head"][1], "preview", filled_type_template(client)).status_code == 403


def test_problems_are_reported_with_their_location(people):
    client = people["hr_head"][1]
    wb = load(filled_type_template(client))
    rules_sheet = wb["Rules"]
    put(rules_sheet, 5, {"rule:priority": 3, "cond:achievement_percent:min": 95,
                         "cond:achievement_percent:min_inclusive": "at least (≥)", "out:score": 5})
    put(rules_sheet, 6, {"rule:priority": 4, "cond:achievement_percent:min": "=1+1",
                         "cond:achievement_percent:min_inclusive": "greater", "out:score": 2})
    put(rules_sheet, 7, {"rule:priority": 5, "cond:actual:min": 1,
                         "cond:actual:min_inclusive": "at least (≥)", "out:score": 1})
    rules_sheet.cell(row=1, column=rules_sheet.max_column + 1, value="cond:bogus:min")
    setting(wb["Matrix"], "name", None)
    preview = send(client, "preview", save(wb)).json()
    found = {(p["location"], p["message"]) for p in preview["problems"]}
    text = " | ".join(f"{loc}: {msg}" for loc, msg in found)
    assert "Contains a formula" in text and "Rules sheet, cell" in text
    assert "'greater' is not one of" in text
    assert "Uses 'actual', which is not marked Use = Yes" in text
    assert "does not match a field on the Fields sheet" in text
    assert "Give the matrix a name." in text
    assert any(loc == "Rules sheet" and "Rule 1 ('Above target') and Rule 3 overlap" in msg
               for loc, msg in found)
    assert not preview["ready"]


def test_confirm_is_all_or_nothing(people, db, monkeypatch):
    client = people["hr_head"][1]
    wb = load(filled_type_template(client))
    put(wb["Rules"], 5, {"rule:priority": 2, "cond:achievement_percent:min": 10,
                         "cond:achievement_percent:min_inclusive": "at least (≥)", "out:score": 1})
    before = counts(db)
    resp = send(client, "confirm", save(wb))
    assert resp.status_code == 422 and counts(db) == before

    def failing(*args, **kwargs):
        raise RuntimeError("audit store unavailable")

    monkeypatch.setattr(matrix_service, "record_event", failing)
    with pytest.raises(RuntimeError):
        send(client, "confirm", filled_type_template(client, "Atomic matrix"))
    assert counts(db) == before


def test_upload_targets_and_scope(people, org):
    client, plant_hr = people["hr_head"][1], people["hr_head_p1"][1]
    created = send(client, "confirm", filled_type_template(client, "Target matrix")).json()
    detail = client.get(f"/api/matrices/{created['matrix_id']}").json()
    data = client.get(f"/api/matrices/{created['matrix_id']}/versions/{created['version_id']}"
                      "/template").content
    assert send(client, "preview", data).json()["action"] == "update_draft"
    client.post(f"/api/matrices/{detail['id']}/versions/{created['version_id']}/submit", json={})
    blocked = send(client, "preview", data).json()
    assert any("awaiting CHRO approval" in p["message"] for p in blocked["problems"])
    # A plant-scoped HR Head can't configure a company-wide matrix.
    own = load(filled_type_template(client, "Plant upload"))
    preview = send(plant_hr, "preview", save(own)).json()
    assert any("within your own scope" in p["message"] for p in preview["problems"])
