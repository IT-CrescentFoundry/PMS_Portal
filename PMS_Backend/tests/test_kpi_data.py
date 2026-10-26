"""Monthly KPI data import: contract, mapping, validation, idempotency, scope, audit."""
from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.models import AuditEvent, KpiDataImport, KpiMonthlyValue
from app.models.enums import ScopeType
from tests.conftest import csv_bytes, import_data, quarter_rows

JULY = [["TST001", "TST-P1", "TST-Q", "TKPI-1", "2026-07", 30],
        ["TST001", "TST-P1", "TST-Q", "TKPI-2", "2026-07", 0],
        ["TST001", "TST-P1", "TST-Q", "TKPI-3", "2026-07", 9]]


def stored(db) -> int:
    return db.scalar(select(func.count()).select_from(KpiMonthlyValue))


def issues(resp) -> list[tuple[int | None, str]]:
    return [(i["row"], i["message"]) for i in resp.json()["issues"]]


def test_preview_writes_nothing_and_confirm_stores_everything(scoring_env, db):
    head = scoring_env.clients["head"]
    preview = import_data(head, scoring_env, "preview", csv_bytes(JULY), "2026-07")
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["ready"] and body["counts"] == {"new": 3, "replace": 0, "unchanged": 0}
    assert body["mapping"]["employee_code"] == "Employee Code" and body["mapping"]["actual"] == "Actual"
    assert stored(db) == 0
    confirmed = import_data(head, scoring_env, "confirm", csv_bytes(JULY), "2026-07")
    assert confirmed.json() == {"import_id": confirmed.json()["import_id"], "created": 3,
                                "replaced": 0, "unchanged": 0, "rows": 3}
    assert stored(db) == 3
    record = db.get(KpiDataImport, confirmed.json()["import_id"])
    assert record.file_name == "data.csv" and len(record.file_sha256) == 64
    assert record.mapping["kpi"] == "KPI"
    value = db.scalars(select(KpiMonthlyValue).where(KpiMonthlyValue.source_row == 2)).one()
    assert value.data_import_id == record.id and str(value.actual_value) == "30.0000"
    event = db.scalars(select(AuditEvent).where(AuditEvent.action == "kpi_data.import")).one()
    assert event.changes["created"] == 3 and event.changes["file"] == "data.csv"


def test_reimport_is_idempotent_and_replacing_needs_confirmation(scoring_env, db):
    head = scoring_env.clients["head"]
    import_data(head, scoring_env, "confirm", csv_bytes(JULY), "2026-07")
    again = import_data(head, scoring_env, "preview", csv_bytes(JULY), "2026-07").json()
    assert again["counts"] == {"new": 0, "replace": 0, "unchanged": 3} and again["ready"]
    assert again["duplicate_import"] is not None
    assert any("already imported" in w for w in again["warnings"])
    changed = [r[:5] + [35] if r[3] == "TKPI-1" else r for r in JULY]
    refused = import_data(head, scoring_env, "confirm", csv_bytes(changed), "2026-07")
    assert refused.status_code == 422
    assert any("Confirm replacing existing values" in m for _r, m in issues(refused))
    replaced = import_data(head, scoring_env, "confirm", csv_bytes(changed), "2026-07",
                           replace_existing=True)
    assert replaced.json()["replaced"] == 1 and replaced.json()["unchanged"] == 2
    value = db.scalars(select(KpiMonthlyValue).where(
        KpiMonthlyValue.kpi_definition_id == scoring_env.env.kpis["profit"].id)).one()
    assert str(value.actual_value) == "35.0000"
    event = db.scalars(select(AuditEvent).where(AuditEvent.action == "kpi_data.import")
                       .order_by(AuditEvent.id.desc())).first()
    assert event.changes["replaced_values"][0]["from"] == "30.0000"


def test_columns_can_be_mapped_explicitly(scoring_env):
    head = scoring_env.clients["head"]
    data = csv_bytes([[r[0], r[1], r[2], r[3], r[5]] for r in JULY],
                     header=["Emp", "Site", "Dept", "Metric", "Value"])
    unmapped = import_data(head, scoring_env, "preview", data, "2026-07")
    assert unmapped.status_code == 422 and "Map the Employee Code" in unmapped.json()["issues"][0]["message"]
    mapping = {"employee_code": "Emp", "plant_code": "Site", "department_code": "Dept",
               "kpi": "Metric", "actual": "Value"}
    mapped = import_data(head, scoring_env, "preview", data, "2026-07", mapping=mapping)
    assert mapped.json()["ready"] and mapped.json()["counts"]["new"] == 3
    bad = import_data(head, scoring_env, "preview", data, "2026-07",
                      mapping={**mapping, "actual": "Emp"})
    assert bad.status_code == 422 and "mapped to more than one field" in bad.json()["issues"][0]["message"]


def test_excel_files_with_dates_and_numbers(scoring_env):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Employee Code", "Plant Code", "Department Code", "KPI", "Period", "Actual"])
    ws.append(["TST001", "TST-P1", "TST-Q", "Test profit", datetime(2026, 7, 1), 30.5])
    buffer = BytesIO()
    wb.save(buffer)
    resp = import_data(scoring_env.clients["head"], scoring_env, "preview", buffer.getvalue(),
                       "2026-07", name="data.xlsx")
    assert resp.json()["ready"], resp.json()["issues"]
    assert resp.json()["lines"][0]["actual"] == "30.5" and resp.json()["lines"][0]["kpi_code"] == "TKPI-1"


@pytest.mark.parametrize("row, expected", [
    (["TST999", "TST-P1", "TST-Q", "TKPI-1", "2026-07", 1], "No published KPI sheet in your scope"),
    (["TST001", "", "", "TKPI-1", "2026-07", 1], "more than one assignment"),
    (["TST001", "TST-P1", "TST-Q", "TKPI-4", "2026-07", 1], "is not on TST001's KPI sheet"),
    (["TST001", "TST-P1", "TST-Q", "TKPI-1", "2026-08", 1], "does not match the month"),
    (["TST001", "TST-P1", "TST-Q", "TKPI-1", "2026-07", ""], "actual value is missing"),
    (["TST001", "TST-P1", "TST-Q", "TKPI-1", "2026-07", "abc"], "is not a number"),
    (["TST001", "TST-P1", "TST-Q", "TKPI-1", "2026-07", "nan"], "is not a number"),
    (["TST001", "TST-P1", "TST-Q", "TKPI-2", "2026-07", "1.5"], "whole number"),
    (["TST002", "TST-P2", "TST-Q", "TKPI-1", "2026-07", 1], "No published KPI sheet"),  # draft
])
def test_row_problems_name_the_row(scoring_env, db, row, expected):
    resp = import_data(scoring_env.clients["head"], scoring_env, "preview", csv_bytes([row]), "2026-07")
    assert any(r == 2 and expected in m for r, m in issues(resp)), issues(resp)
    assert stored(db) == 0


def test_units_and_duplicates(scoring_env):
    head = scoring_env.clients["head"]
    data = csv_bytes([["TST001", "TST-P1", "TST-Q", "TKPI-1", "2026-07", 30, "kg"]],
                     header=["Employee Code", "Plant Code", "Department Code", "KPI", "Period",
                             "Actual", "Unit"])
    assert any("does not match the KPI sheet's unit" in m for _r, m in
               issues(import_data(head, scoring_env, "preview", data, "2026-07")))
    twice = csv_bytes([JULY[0], JULY[0]])
    assert any(r == 3 and "already given in row 2" in m for r, m in
               issues(import_data(head, scoring_env, "preview", twice, "2026-07")))


@pytest.mark.parametrize("period, expected", [
    ("2026-10", "is not a month of TQ1"),
    ("July", "YYYY-MM"),
])
def test_the_month_must_belong_to_the_quarter(scoring_env, period, expected):
    resp = import_data(scoring_env.clients["head"], scoring_env, "preview", csv_bytes(JULY), period)
    assert resp.status_code == 422 and expected in resp.json()["issues"][0]["message"]


def test_the_month_must_have_started(scoring_env, monkeypatch):
    from datetime import date

    from app.services import clock

    monkeypatch.setattr(clock, "today", lambda: date(2026, 8, 10))
    rows = [[*r[:4], "2026-09", r[5]] for r in JULY]
    resp = import_data(scoring_env.clients["head"], scoring_env, "preview", csv_bytes(rows), "2026-09")
    assert resp.status_code == 422 and "September 2026 has not started yet" in resp.json()["issues"][0]["message"]


def test_cross_plant_rows_are_out_of_scope(scoring_env, as_role):
    o = scoring_env.env.org
    _, p2_head = as_role("tst.data.p2", [("DEPT_HEAD", [(ScopeType.PLANT_DEPARTMENT, o.plants["P2"],
                                                         o.departments["Q"])])])
    resp = import_data(p2_head, scoring_env, "preview", csv_bytes(JULY[:1]), "2026-07")
    assert any("No published KPI sheet in your scope" in m for _r, m in issues(resp))
    _, reader = as_role("tst.data.reader", [("BUDGET_USER", [(ScopeType.GLOBAL, None, None)])])
    assert import_data(reader, scoring_env, "preview", csv_bytes(JULY), "2026-07").status_code == 403


def test_confirm_with_any_problem_writes_nothing(scoring_env, db):
    rows = [*JULY, ["TST001", "TST-P1", "TST-Q", "TKPI-1", "2026-08", 5]]
    resp = import_data(scoring_env.clients["head"], scoring_env, "confirm", csv_bytes(rows), "2026-07")
    assert resp.status_code == 422 and stored(db) == 0
    assert db.scalar(select(func.count()).select_from(KpiDataImport)) == 0


def test_unreadable_files(scoring_env):
    head = scoring_env.clients["head"]
    assert import_data(head, scoring_env, "preview", b"", "2026-07").status_code == 422
    resp = import_data(head, scoring_env, "preview", b"PK\x03\x04junk", "2026-07", name="data.xlsx")
    assert resp.status_code == 422 and "could not be read" in resp.json()["issues"][0]["message"]
    resp = import_data(head, scoring_env, "preview", csv_bytes(JULY), "2026-07", name="data.txt")
    assert "CSV (.csv) or Excel (.xlsx)" in resp.json()["issues"][0]["message"]
    latin = "Employee Code,KPI\nTSTé,TKPI-1\n".encode("latin-1")
    resp = import_data(head, scoring_env, "preview", latin, "2026-07")
    assert "UTF-8" in resp.json()["issues"][0]["message"]


def test_template_lists_the_kpis_needing_data(scoring_env):
    resp = scoring_env.clients["head"].get("/api/kpi/data/template", params={
        "quarter_id": scoring_env.env.q1.id, "period": "2026-08"})
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.content), read_only=True)
    rows = list(wb["Data"].iter_rows(values_only=True))
    assert rows[0][:6] == ("Employee Code", "Plant Code", "Department Code", "KPI", "Period", "Actual")
    data_rows = [r for r in rows[1:] if r[0]]
    assert {(r[1], r[3]) for r in data_rows} == {
        ("TST-P1", "TKPI-1"), ("TST-P1", "TKPI-2"), ("TST-P1", "TKPI-3"), ("TST-P2", "TKPI-1")}
    assert all(r[4] == "2026-08" for r in data_rows)
    # The filled-in template is accepted as it is.
    filled = csv_bytes([[*r[:5], 1] for r in data_rows])
    resp = import_data(scoring_env.clients["head"], scoring_env, "preview", filled, "2026-08")
    assert resp.json()["ready"], resp.json()["issues"]


def test_imports_are_listed_with_file_metadata_only(scoring_env):
    head = scoring_env.clients["head"]
    import_data(head, scoring_env, "confirm", csv_bytes(quarter_rows({"TKPI-1": [30]})), "2026-07")
    listed = head.get("/api/kpi/data/imports", params={"quarter_id": scoring_env.env.q1.id}).json()
    assert listed[0]["file_name"] == "data.csv" and listed[0]["rows_created"] == 1
    assert set(listed[0]) == {"id", "period", "file_name", "file_format", "file_size", "row_count",
                              "rows_created", "rows_replaced", "rows_unchanged", "created_by",
                              "created_at"}
