"""Department Head Excel uploads: parsing, row validation, scope, all-or-nothing confirm."""
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.models import AuditEvent, KpiSheet
from app.models.enums import ScopeType

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADERS = ["Employee Code *", "Plant Code", "Department Code", "KPI *", "Weightage %", "Target per Quarter"]


def workbook(rows, headers=HEADERS, percent=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "KPIs"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    if percent:
        for cell in ws["E"][1:]:
            cell.number_format = "0%"
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def send(client, step, quarter, data, filename="kpis.xlsx"):
    return client.post(f"/api/kpi/uploads/{step}", data={"quarter_id": str(quarter.id)},
                       files={"file": (filename, data, XLSX)})


# TST001 works at P1/Q and P2/Q, so its rows name the unit; TST002 and TST003 are at P2/Q only.
VALID = [
    ["TST001", "TST-P1", "TST-Q", "Test profit", 40, 120],
    ["TST001", "TST-P1", "TST-Q", "TKPI-4", 35, 95],          # by code
    ["TST001", "TST-P1", "TST-Q", "Test claims", "25%", None],
    ["TST002", None, None, "test  PROFIT", 60, 100],           # spacing and case are ignored
    ["TST002", None, None, "Test inventory days", 40, 30],
]


@pytest.fixture
def qa_head(as_role, kpi_env):
    return as_role("tst.up.qa", [("DEPT_HEAD", [(ScopeType.DEPARTMENT, None, kpi_env.org.departments["Q"])])])


def sheet_count(db):
    return db.scalar(select(func.count()).select_from(KpiSheet))


def test_preview_then_confirm_creates_draft_sheets(qa_head, kpi_env, db):
    client = qa_head[1]
    data = workbook(VALID)
    preview = send(client, "preview", kpi_env.q1, data)
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["issues"] == [] and body["row_count"] == 5
    assert [(s["employee"]["employee_code"], s["plant"]["code"], s["action"]) for s in body["sheets"]] == [
        ("TST001", "TST-P1", "create"), ("TST002", "TST-P2", "create")]
    assert [i["weightage"] for i in body["sheets"][0]["items"]] == [40.0, 35.0, 25.0]
    assert sheet_count(db) == 0                                       # preview writes nothing

    confirmed = send(client, "confirm", kpi_env.q1, data)
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json() == {"sheets_created": 2, "sheets_updated": 0, "items": 5, "rows": 5}
    sheets = db.scalars(select(KpiSheet)).all()
    assert {s.status.value for s in sheets} == {"DRAFT"}
    uploads = db.scalars(select(AuditEvent.changes).where(AuditEvent.action == "kpi_sheet.upload")).all()
    assert len(uploads) == 2 and all(c["file"] == "kpis.xlsx" for c in uploads)

    # Uploading again replaces the draft items rather than adding to them.
    again = send(client, "preview", kpi_env.q1, data).json()
    assert {s["action"] for s in again["sheets"]} == {"replace"}


def test_percent_formatted_cells_are_read_as_percentages(qa_head, kpi_env):
    data = workbook([["TST002", None, None, "Test profit", 0.6, 100],
                     ["TST002", None, None, "Test inventory days", 0.4, 30]], percent=True)
    body = send(qa_head[1], "preview", kpi_env.q1, data).json()
    assert body["issues"] == [] and [i["weightage"] for i in body["sheets"][0]["items"]] == [60.0, 40.0]


def test_formula_float_noise_is_tolerated_but_three_decimals_are_not(qa_head, kpi_env):
    noisy = workbook([["TST002", None, None, "Test profit", 0.1 + 0.2, 100],      # 0.30000000000000004
                      ["TST002", None, None, "Test inventory days", 0.7, 30]], percent=True)
    assert send(qa_head[1], "preview", kpi_env.q1, noisy).json()["issues"] == []
    precise = workbook([["TST002", None, None, "Test profit", 33.333, 100],
                        ["TST002", None, None, "Test inventory days", 66.667, 30]])
    messages = [i["message"] for i in send(qa_head[1], "preview", kpi_env.q1, precise).json()["issues"]]
    assert any("two decimal places" in m for m in messages)


@pytest.mark.parametrize("row, expected", [
    (["TST999", None, None, "Test profit", 100, 10], "TST999 not found"),
    (["TST002", None, None, "No such KPI", 100, 10], "not in the KPI library"),
    (["TST001", None, None, "Test profit", 100, 10], "more than one unit"),
    (["TST002", "TST-P9", None, "Test profit", 100, 10], "Plant code 'TST-P9' not found"),
    (["TST002", None, None, "Test profit", "abc", 10], "Weightage must be a number"),
    (["TST002", None, None, "Test profit", 100, "nan"], "Target per Quarter must be a number"),
    (["TST002", None, None, "Test profit", 100, 10 ** 15], "below 100 trillion"),
    (["TST002", None, None, "Test yield P1", 100, 10], "does not apply"),
    (["TST002", None, None, "Test profit", 90, 10], "must total exactly 100"),
    (["TST004", None, None, "Test maintenance", 100, 10], "no assignment in your scope"),  # Maintenance
])
def test_row_problems_are_reported_with_row_numbers(qa_head, kpi_env, row, expected):
    body = send(qa_head[1], "preview", kpi_env.q1, workbook([row])).json()
    assert any(expected in i["message"] and i["row"] == 2 for i in body["issues"]), body["issues"]


def test_duplicate_kpi_rows_are_detected(qa_head, kpi_env):
    rows = [["TST002", None, None, "Test profit", 50, 10], ["TST002", None, None, "TKPI-1", 50, 10]]
    issues = send(qa_head[1], "preview", kpi_env.q1, workbook(rows)).json()["issues"]
    assert any("appears more than once" in i["message"] and i["row"] == 3 for i in issues)


def test_confirm_with_any_problem_writes_nothing(qa_head, kpi_env, db):
    rows = VALID + [["TST003", None, None, "Test profit", 70, 10]]   # TST003 totals 70 %
    resp = send(qa_head[1], "confirm", kpi_env.q1, workbook(rows))
    assert resp.status_code == 422 and any("TST003" in i["message"] for i in resp.json()["issues"])
    assert sheet_count(db) == 0


@pytest.mark.parametrize("data, filename, expected", [
    (b"", "kpis.xlsx", "empty"),
    (b"name,weight\nx,1\n", "kpis.xlsx", "Upload an Excel workbook"),
    (b"PK\x03\x04 not really a zip", "kpis.xlsx", "could not be read"),
    (None, "kpis.csv", "Upload an Excel workbook"),
    (b"PK\x03\x04" + b"0" * 1_000_001, "kpis.xlsx", "larger than 1 MB"),
], ids=["empty", "csv-content", "broken-zip", "csv-name", "too-large"])
def test_unreadable_files_are_refused(qa_head, kpi_env, data, filename, expected):
    resp = send(qa_head[1], "preview", kpi_env.q1, workbook(VALID) if data is None else data, filename)
    assert resp.status_code == 422 and expected in resp.json()["issues"][0]["message"]


def test_missing_required_columns(qa_head, kpi_env):
    data = workbook([["TST002", "Test profit"]], headers=["Employee Code", "KPI"])
    resp = send(qa_head[1], "preview", kpi_env.q1, data)
    assert resp.status_code == 422 and "Weightage %" in resp.json()["issues"][0]["message"]


def test_out_of_scope_rows_cannot_be_written(as_role, kpi_env, db):
    o = kpi_env.org
    _, p2_head = as_role("tst.up.p2", [("DEPT_HEAD", [(ScopeType.PLANT_DEPARTMENT, o.plants["P2"],
                                                        o.departments["Q"])])])
    rows = [["TST001", "TST-P1", "TST-Q", "Test profit", 100, 10]]           # P1 is not theirs
    resp = send(p2_head, "confirm", kpi_env.q1, workbook(rows))
    assert resp.status_code == 422 and "no assignment in your scope" in resp.json()["issues"][0]["message"]
    assert sheet_count(db) == 0

    _, reader = as_role("tst.up.reader", [("BUDGET_USER", [(ScopeType.GLOBAL, None, None)])])
    assert send(reader, "preview", kpi_env.q1, workbook(VALID)).status_code == 403


def test_published_sheets_and_started_quarters_are_protected(qa_head, kpi_env, db, as_role):
    client = qa_head[1]
    send(client, "confirm", kpi_env.q1, workbook(VALID))
    sheet_id = db.scalar(select(KpiSheet.id).join(KpiSheet.assignment).where(KpiSheet.cycle_quarter_id == kpi_env.q1.id)
                         .order_by(KpiSheet.id))
    assert client.post(f"/api/kpi/sheets/{sheet_id}/publish").status_code == 200
    issues = send(client, "preview", kpi_env.q1, workbook(VALID)).json()["issues"]
    assert any("is published; reopen it" in i["message"] for i in issues)

    _, chro = as_role("tst.up.chro", [("CHRO", [(ScopeType.GLOBAL, None, None)])])
    chro.post(f"/api/kpi/quarters/{kpi_env.q1.id}/start-scoring")
    resp = send(client, "confirm", kpi_env.q1, workbook(VALID))
    assert resp.status_code == 422 and "Scoring has started" in resp.json()["issues"][0]["message"]


def test_template_lists_only_visible_kpis(as_role, kpi_env):
    o = kpi_env.org
    _, p2_head = as_role("tst.up.tpl", [("DEPT_HEAD", [(ScopeType.PLANT_DEPARTMENT, o.plants["P2"],
                                                         o.departments["Q"])])])
    resp = p2_head.get("/api/kpi/uploads/template")
    assert resp.status_code == 200 and resp.headers["content-type"] == XLSX
    wb = load_workbook(BytesIO(resp.content), read_only=True)
    assert wb.sheetnames == ["KPIs", "Instructions", "KPI Library", "Scoring Policies"]
    names = {row[1] for row in wb["KPI Library"].iter_rows(min_row=2, values_only=True)}
    assert {"Test profit", "Test claims", "Test inventory days"} <= names
    assert not names & {"Test yield P1", "Test maintenance"}
    headers = [c for c in next(wb["KPIs"].iter_rows(max_row=1, values_only=True))]
    assert headers[0] == "Employee Code *" and "Weightage % *" in headers


def test_filled_in_template_uploads_cleanly(qa_head, kpi_env):
    client = qa_head[1]
    wb = load_workbook(BytesIO(client.get("/api/kpi/uploads/template").content))
    ws = wb["KPIs"]
    ws.append(["TST002", None, None, "Test profit", 60, 100, "%", "TPOL-A", "Synthetic note"])
    ws.append(["TST002", None, None, "Test inventory days", 40, 30])
    buffer = BytesIO()
    wb.save(buffer)
    body = send(client, "preview", kpi_env.q1, buffer.getvalue()).json()
    assert body["issues"] == []
    assert body["sheets"][0]["items"][0]["scoring_policy"] == "TPOL-A"
