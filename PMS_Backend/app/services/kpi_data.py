"""Monthly KPI data import (CSV or XLSX): validate-only preview, then confirm.

Import contract (also in docs/api.md and in the generated template):
  Employee Code *     the employee's code in the PMS
  Plant Code          } needed only when the employee has several
  Department Code     } assignments with a published sheet in the month
  KPI *               the KPI's code or name; it must be on the employee's
                      published KPI sheet for the quarter
  Period              YYYY-MM; optional, must equal the month chosen
  Actual              the month's value in the KPI's unit; required when the
                      KPI's scoring matrix uses the actual
  Unit                optional; must match the KPI sheet's unit
  <input label>       one column per extra matrix input imported with the
                      data (e.g. an audit grade), declared by the frozen
                      scoring matrix of the KPI
Columns are matched to these fields by name (common variants are
recognised) or by an explicit mapping chosen in the portal.

Rules:
  * Only months that have started, inside the chosen quarter.
  * Only PUBLISHED or LOCKED sheets in the importer's KPI-writing scope, and
    never after a sheet's scores are published.
  * An identical value already stored is left unchanged; a different value
    replaces it only when the import is confirmed with replace_existing. A
    file identical to an earlier import of the same month is flagged.
  * Preview writes nothing. Confirm re-validates and writes everything in one
    transaction, or nothing. The file itself is never stored: only its name,
    size, SHA-256 and column mapping, and each value's source row.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload, selectinload

from app.auth.permissions import Permission, Principal
from app.models import (
    EmployeeAssignment,
    KpiDataImport,
    KpiMonthlyValue,
    KpiSheet,
    KpiSheetItem,
    KpiSheetResult,
    KpiSheetStatus,
    MatrixVersion,
    PlantDepartment,
    PmsCycleQuarter,
    ScoreResultStatus,
    User,
)
from app.services import clock
from app.services import matrix_rules as rules
from app.services import matrix_service
from app.services.audit import AuditActor, record_event
from app.services.errors import Issue, PermissionDeniedError, ValidationFailedError
from app.services.kpi_sheets import COLLECTION_STATUSES, QUARTER_LABELS, get_quarter
from app.services.matrix_types import ACTUAL_SOURCES, DataType, Source

MAX_BYTES = 2_000_000
MAX_ROWS = 5000
MAX_ACTUAL = Decimal("1e14")  # Numeric(18, 4)
VALUE_PLACES = Decimal("0.0001")
ACCEPTED_SHEETS = (KpiSheetStatus.PUBLISHED, KpiSheetStatus.LOCKED)


@dataclass(frozen=True)
class ContractField:
    key: str
    label: str
    required: bool
    description: str
    aliases: tuple[str, ...] = ()


CONTRACT: tuple[ContractField, ...] = (
    ContractField("employee_code", "Employee Code", True, "The employee's code in the PMS.",
                  ("employeecode", "empcode", "employeeno", "employee", "empno")),
    ContractField("plant_code", "Plant Code", False,
                  "Needed when the employee has more than one assignment.", ("plantcode", "plant")),
    ContractField("department_code", "Department Code", False,
                  "Needed when the employee has more than one assignment.",
                  ("departmentcode", "deptcode", "department", "dept")),
    ContractField("kpi", "KPI", True, "The KPI's code or name, as on the KPI sheet.",
                  ("kpi", "kpicode", "kpiname", "kra", "kracode")),
    ContractField("period", "Period", False, "YYYY-MM; must equal the month being imported.",
                  ("period", "month", "periodyyyymm")),
    ContractField("actual", "Actual", False,
                  "The month's value in the KPI's unit (required when the scoring matrix "
                  "uses the actual).", ("actual", "actualvalue", "value", "monthlyactual")),
    ContractField("unit", "Unit", False, "Optional; must match the KPI sheet's unit.",
                  ("unit", "uom")),
)


def _norm_header(text: Any) -> str:
    return re.sub(r"[\s_\-()*./]", "", str(text or "")).lower()


def _norm(text: Any) -> str:
    return " ".join(str(text).split()).casefold() if text is not None else ""


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value: Any) -> str | None:
    if _blank(value):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_period(text: str) -> date:
    match = re.fullmatch(r"\s*(\d{4})-(\d{1,2})\s*", text or "")
    if not match or not 1 <= int(match.group(2)) <= 12:
        raise ValidationFailedError([Issue("Choose the month as YYYY-MM.", "period")])
    return date(int(match.group(1)), int(match.group(2)), 1)


# --- reading files ---------------------------------------------------------------------

@dataclass
class Table:
    file_format: str
    headers: list[str]
    rows: list[tuple[int, dict[str, Any]]]  # file row number, header -> value


def read_table(file_name: str, data: bytes) -> Table:
    if not data:
        raise ValidationFailedError([Issue("The file is empty.")])
    if len(data) > MAX_BYTES:
        raise ValidationFailedError([Issue("The file is larger than 2 MB.")])
    name = (file_name or "").lower()
    raw_rows: list[list[Any]]
    if name.endswith(".csv"):
        file_format = "csv"
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValidationFailedError([Issue("Save the CSV file with UTF-8 encoding.")]) from None
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        raw_rows = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    elif name.endswith(".xlsx"):
        file_format = "xlsx"
        if not data.startswith(b"PK\x03\x04"):
            raise ValidationFailedError([Issue("The file is not a valid Excel workbook (.xlsx).")])
        try:
            # Cached values of formula cells are used; nothing is evaluated here.
            wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        except (InvalidFileException, BadZipFile, KeyError, OSError, ValueError):
            raise ValidationFailedError([Issue("The file could not be read as an Excel "
                                               "workbook (.xlsx).")]) from None
        try:
            ws = wb["Data"] if "Data" in wb.sheetnames else wb.worksheets[0]
            raw_rows = []
            for row in ws.iter_rows(min_row=1, values_only=True):
                if len(raw_rows) > MAX_ROWS + 50:
                    break
                raw_rows.append(list(row))
        finally:
            wb.close()
    else:
        raise ValidationFailedError([Issue("Upload a CSV (.csv) or Excel (.xlsx) file.")])

    header_index = next((i for i, r in enumerate(raw_rows) if any(not _blank(v) for v in r)), None)
    if header_index is None:
        raise ValidationFailedError([Issue("The file has no header row.")])
    headers = [(_text(v) or "") for v in raw_rows[header_index]]
    seen: set[str] = set()
    for h in headers:
        if h and _norm_header(h) in seen:
            raise ValidationFailedError([Issue(f"The column '{h}' appears more than once.")])
        if h:
            seen.add(_norm_header(h))
    rows: list[tuple[int, dict[str, Any]]] = []
    for offset, values in enumerate(raw_rows[header_index + 1:], header_index + 2):
        record = {h: (v.strip() if isinstance(v, str) else v)
                  for h, v in zip(headers, values) if h and not _blank(v)}
        if record:
            if len(rows) >= MAX_ROWS:
                raise ValidationFailedError([Issue(f"The file has more than {MAX_ROWS} rows.")])
            rows.append((offset, record))
    return Table(file_format, [h for h in headers if h], rows)


# --- context: sheets and items in scope ----------------------------------------------------

@dataclass
class ItemContext:
    sheet: KpiSheet
    item: KpiSheetItem
    definition: rules.Definition | None  # frozen matrix version; None for Stage J items
    problem: str | None  # why the item can't take data (e.g. altered configuration)

    @property
    def needs_actual(self) -> bool:
        if self.definition is None:
            return True
        return any(f.source in ACTUAL_SOURCES for f in self.definition.inputs)

    @property
    def whole_actual(self) -> bool:
        return self.definition is not None and any(
            f.source == Source.ACTUAL and f.data_type == DataType.INTEGER
            for f in self.definition.inputs)

    @property
    def imported(self) -> list[rules.FieldDef]:
        if self.definition is None:
            return []
        return [f for f in self.definition.inputs if f.source == Source.IMPORTED]


def _item_context(sheet: KpiSheet, item: KpiSheetItem) -> ItemContext:
    if item.matrix_version is None:
        return ItemContext(sheet, item, None, None)
    try:
        return ItemContext(sheet, item, matrix_service.frozen_definition(item.matrix_version), None)
    except (rules.EvaluationError, matrix_service.IntegrityProblem) as exc:
        return ItemContext(sheet, item, None, str(exc))


def sheets_in_scope(db: Session, principal: Principal, quarter: PmsCycleQuarter,
                    permission: Permission = Permission.KPI_WRITE) -> list[KpiSheet]:
    scope = principal.scope_for(permission)
    if scope.is_empty:
        return []
    return list(db.scalars(
        select(KpiSheet)
        .join(KpiSheet.assignment).join(EmployeeAssignment.plant_department)
        .where(KpiSheet.cycle_quarter_id == quarter.id, KpiSheet.status.in_(ACCEPTED_SHEETS),
               scope.unit_condition())
        .options(
            joinedload(KpiSheet.assignment).joinedload(EmployeeAssignment.employee),
            joinedload(KpiSheet.assignment).joinedload(EmployeeAssignment.plant_department)
            .joinedload(PlantDepartment.plant),
            joinedload(KpiSheet.assignment).joinedload(EmployeeAssignment.plant_department)
            .joinedload(PlantDepartment.department),
            selectinload(KpiSheet.items).joinedload(KpiSheetItem.kpi_definition),
            selectinload(KpiSheet.items).joinedload(KpiSheetItem.matrix_version)
            .joinedload(MatrixVersion.matrix),
        )
    ).unique().all())


def _active_in(sheet: KpiSheet, month: date) -> bool:
    a = sheet.assignment
    return a.effective_from <= clock.month_end(month) and (
        a.effective_to is None or a.effective_to >= month)


def extra_fields(contexts: Sequence[ItemContext]) -> list[ContractField]:
    """Contract fields for the extra matrix inputs imported with the data."""
    found: dict[str, ContractField] = {}
    for ctx in contexts:
        for f in ctx.imported:
            found.setdefault(f.key, ContractField(
                f"input:{f.key}", f.label, False,
                f"Extra input for the scoring matrix ({f.data_type.value.lower()}"
                + (f"; one of {', '.join(f.allowed_values)}" if f.allowed_values else "") + ").",
                (_norm_header(f.key), _norm_header(f.label))))
    return sorted(found.values(), key=lambda c: c.label)


# --- checking an import --------------------------------------------------------------------

@dataclass
class Line:
    row: int
    sheet_id: int
    kpi_definition_id: int
    employee_code: str
    employee_name: str
    unit_label: str
    kpi_code: str
    kpi_name: str
    actual: Decimal | None
    inputs: dict[str, Any] | None
    action: str  # new | replace | unchanged
    previous_actual: Decimal | None = None
    previous_inputs: dict[str, Any] | None = None
    note: str | None = None


@dataclass
class ImportCheck:
    quarter: PmsCycleQuarter
    period: date
    file_name: str
    file_format: str
    size: int
    sha256: str
    columns: list[str]
    contract: list[ContractField]
    mapping: dict[str, str | None]
    lines: list[Line] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    duplicate_of: KpiDataImport | None = None
    row_count: int = 0

    @property
    def ready(self) -> bool:
        return not self.issues and bool(self.lines)

    def count(self, action: str) -> int:
        return sum(1 for line in self.lines if line.action == action)


def auto_mapping(contract: Sequence[ContractField], headers: Sequence[str]) -> dict[str, str | None]:
    by_norm = {_norm_header(h): h for h in headers}
    mapping: dict[str, str | None] = {}
    for c in contract:
        candidates = (_norm_header(c.label), _norm_header(c.key), *c.aliases)
        mapping[c.key] = next((by_norm[n] for n in candidates if n in by_norm), None)
    return mapping


def _apply_mapping(contract: Sequence[ContractField], headers: Sequence[str],
                   override: Mapping[str, str | None] | None) -> dict[str, str | None]:
    mapping = auto_mapping(contract, headers)
    if override:
        known = {c.key for c in contract}
        unknown = sorted(set(override) - known)
        if unknown:
            raise ValidationFailedError([Issue(f"Unknown field(s) in the mapping: "
                                               f"{', '.join(unknown)}.", "mapping")])
        for key, header in override.items():
            if header not in (None, "") and header not in headers:
                raise ValidationFailedError([Issue(f"The file has no column '{header}'.",
                                                   "mapping")])
            mapping[key] = header or None
    used = [h for h in mapping.values() if h]
    duplicates = sorted({h for h in used if used.count(h) > 1})
    if duplicates:
        raise ValidationFailedError([Issue(f"The column '{duplicates[0]}' is mapped to more than "
                                           "one field.", "mapping")])
    missing = [c.label for c in contract if c.required and not mapping.get(c.key)]
    if missing:
        raise ValidationFailedError([Issue(f"Map the {' and '.join(missing)} column(s).",
                                           "mapping")])
    return mapping


def _parse_actual(raw: Any, ctx: ItemContext) -> tuple[Decimal, str | None]:
    number = rules.parse_decimal(raw)
    if ctx.whole_actual and number != number.to_integral_value():
        raise ValueError("The scoring matrix counts whole numbers; enter a whole number.")
    if abs(number) >= MAX_ACTUAL:
        raise ValueError("The value is too large.")
    rounded = number.quantize(VALUE_PLACES, rounding=ROUND_HALF_UP)
    note = f"Rounded to {rules.decimal_text(rounded)}." if rounded != number else None
    return rounded, note


def check_import(db: Session, principal: Principal, *, quarter_id: int, period: str,
                 file_name: str, data: bytes, mapping: Mapping[str, str | None] | None = None,
                 replace_existing: bool = False) -> ImportCheck:
    """Everything the preview shows; writes nothing."""
    if principal.scope_for(Permission.KPI_WRITE).is_empty:
        raise PermissionDeniedError("You cannot import KPI data.")
    quarter = get_quarter(db, quarter_id)
    month = parse_period(period)
    if quarter.status not in COLLECTION_STATUSES:
        raise ValidationFailedError([Issue(
            f"Data collection for {quarter.name} is closed ({QUARTER_LABELS[quarter.status]}); "
            "a correction needs the quarter returned to data collection by HR.", "quarter_id")])
    if month not in clock.months_between(quarter.start_date, quarter.end_date):
        raise ValidationFailedError([Issue(f"{month:%Y-%m} is not a month of {quarter.name}.",
                                           "period")])
    if month > clock.month_start(clock.today()):
        raise ValidationFailedError([Issue(f"{month:%B %Y} has not started yet.", "period")])

    table = read_table(file_name, data)
    sheets = [s for s in sheets_in_scope(db, principal, quarter) if _active_in(s, month)]
    contexts = {(s.id, i.kpi_definition_id): _item_context(s, i) for s in sheets for i in s.items}
    contract = [*CONTRACT, *extra_fields(list(contexts.values()))]
    check = ImportCheck(quarter=quarter, period=month, file_name=file_name,
                        file_format=table.file_format, size=len(data),
                        sha256=hashlib.sha256(data).hexdigest(), columns=table.headers,
                        contract=contract, mapping=_apply_mapping(contract, table.headers, mapping),
                        row_count=len(table.rows))
    if not table.rows:
        check.issues.append(Issue("The file has no data rows below the header."))
        return check

    by_employee: dict[str, list[KpiSheet]] = {}
    for s in sheets:
        by_employee.setdefault(_norm(s.assignment.employee.employee_code), []).append(s)
    published_results = set(db.scalars(
        select(KpiSheetResult.sheet_id).where(
            KpiSheetResult.sheet_id.in_([s.id for s in sheets] or [0]),
            KpiSheetResult.status == ScoreResultStatus.PUBLISHED)).all())
    existing = {
        (v.sheet_id, v.kpi_definition_id): v
        for v in db.scalars(select(KpiMonthlyValue).where(
            KpiMonthlyValue.sheet_id.in_([s.id for s in sheets] or [0]),
            KpiMonthlyValue.period_month == month)).all()
    }
    seen: dict[tuple[int, int], int] = {}

    def get(values: Mapping[str, Any], key: str) -> Any:
        header = check.mapping.get(key)
        return values.get(header) if header else None

    for row, values in table.rows:
        def issue(message: str, field_key: str | None = None) -> None:
            check.issues.append(Issue(message, field_key, row))

        code = _text(get(values, "employee_code"))
        if not code:
            issue("Employee Code is required.", "employee_code")
            continue
        candidates = by_employee.get(_norm(code), [])
        plant_code = _text(get(values, "plant_code"))
        dept_code = _text(get(values, "department_code"))
        if plant_code:
            candidates = [s for s in candidates
                          if _norm(s.assignment.plant_department.plant.code) == _norm(plant_code)]
        if dept_code:
            candidates = [s for s in candidates if _norm(
                s.assignment.plant_department.department.code) == _norm(dept_code)]
        if not candidates:
            issue(f"No published KPI sheet in your scope for employee {code} in {quarter.name} "
                  f"during {month:%B %Y}" + (" with that plant and department." if plant_code
                                             or dept_code else "."), "employee_code")
            continue
        if len(candidates) > 1:
            issue(f"{code} has more than one assignment with a KPI sheet; fill in Plant Code and "
                  "Department Code.", "plant_code")
            continue
        sheet = candidates[0]
        kpi_text = _text(get(values, "kpi"))
        if not kpi_text:
            issue("KPI is required.", "kpi")
            continue
        item = next((i for i in sheet.items if _norm(kpi_text) in (
            _norm(i.kpi_definition.code), _norm(i.kpi_definition.name))), None)
        if item is None:
            issue(f"'{kpi_text}' is not on {code}'s KPI sheet for {quarter.name}.", "kpi")
            continue
        ctx = contexts[(sheet.id, item.kpi_definition_id)]
        label = f"{code} · {item.kpi_definition.name}"
        if sheet.id in published_results:
            issue(f"{label}: scores for this sheet are published; its data can no longer change.")
            continue
        if ctx.problem:
            issue(f"{label}: {ctx.problem}")
            continue
        period_value = get(values, "period")
        if not _blank(period_value):
            if isinstance(period_value, datetime):  # an Excel date cell
                period_month = clock.month_start(period_value.date())
            elif isinstance(period_value, date):
                period_month = clock.month_start(period_value)
            else:
                try:
                    period_month = parse_period(str(period_value))
                except ValidationFailedError:
                    issue(f"Period '{period_value}' is not YYYY-MM.", "period")
                    continue
            if period_month != month:
                issue(f"Period {period_month:%Y-%m} does not match the month being imported "
                      f"({month:%Y-%m}).", "period")
                continue
        unit = _text(get(values, "unit"))
        if unit and _norm(unit) != _norm(item.unit):
            issue(f"{label}: the unit '{unit}' does not match the KPI sheet's unit "
                  f"'{item.unit}'.", "unit")
            continue
        actual, note = None, None
        raw_actual = get(values, "actual")
        if _blank(raw_actual):
            if ctx.needs_actual:
                issue(f"{label}: the actual value is missing.", "actual")
                continue
        else:
            try:
                actual, note = _parse_actual(raw_actual, ctx)
            except ValueError as exc:
                issue(f"{label}: actual — {exc}", "actual")
                continue
        inputs: dict[str, Any] = {}
        bad = False
        for f in ctx.imported:
            raw = get(values, f"input:{f.key}")
            if _blank(raw):
                if f.required:
                    issue(f"{label}: {f.label} is missing.", f"input:{f.key}")
                    bad = True
                continue
            try:
                inputs[f.key] = rules.jsonable(rules.parse_value(f, raw))
            except ValueError as exc:
                issue(f"{label}: {f.label} — {exc}", f"input:{f.key}")
                bad = True
        if bad:
            continue
        key = (sheet.id, item.kpi_definition_id)
        if key in seen:
            issue(f"{label}: already given in row {seen[key]}.", "kpi")
            continue
        seen[key] = row
        previous = existing.get(key)
        action = "new"
        if previous is not None:
            same = (previous.actual_value == actual
                    and (previous.input_values or {}) == (inputs or {}))
            action = "unchanged" if same else "replace"
            if action == "replace" and not replace_existing:
                issue(f"{label}: a different value for {month:%B %Y} is already stored "
                      f"({rules.decimal_text(previous.actual_value) if previous.actual_value is not None else 'no actual'}). "
                      "Confirm replacing existing values to overwrite it.", "actual")
        unit_ref = sheet.assignment.plant_department
        check.lines.append(Line(
            row=row, sheet_id=sheet.id, kpi_definition_id=item.kpi_definition_id,
            employee_code=sheet.assignment.employee.employee_code,
            employee_name=sheet.assignment.employee.full_name,
            unit_label=f"{unit_ref.department.name} — {unit_ref.plant.name}",
            kpi_code=item.kpi_definition.code, kpi_name=item.kpi_definition.name, actual=actual,
            inputs=inputs or None, action=action,
            previous_actual=previous.actual_value if previous is not None else None,
            previous_inputs=previous.input_values if previous is not None else None, note=note))

    check.duplicate_of = db.scalars(
        select(KpiDataImport).where(KpiDataImport.file_sha256 == check.sha256,
                                    KpiDataImport.period_month == month,
                                    KpiDataImport.cycle_quarter_id == quarter.id)
        .order_by(KpiDataImport.id.desc())).first()
    if check.duplicate_of is not None:
        check.warnings.append(f"This exact file was already imported for {month:%B %Y} on "
                              f"{check.duplicate_of.created_at:%Y-%m-%d %H:%M}. Unchanged values "
                              "are skipped.")
    missing = len(set(contexts) - set(seen) - set(existing))
    if missing:
        check.warnings.append(f"{missing} KPI(s) on published sheets in your scope still have no "
                              f"value for {month:%B %Y}.")
    return check


def confirm_import(db: Session, principal: Principal, actor: AuditActor, *, quarter_id: int,
                   period: str, file_name: str, data: bytes,
                   mapping: Mapping[str, str | None] | None = None,
                   replace_existing: bool = False) -> KpiDataImport:
    """Re-validate and store every value, or none (the caller commits once)."""
    check = check_import(db, principal, quarter_id=quarter_id, period=period, file_name=file_name,
                         data=data, mapping=mapping, replace_existing=replace_existing)
    if check.issues:
        raise ValidationFailedError(check.issues)
    if not check.lines:
        raise ValidationFailedError([Issue("The file has no data rows.")])
    record = KpiDataImport(
        cycle_quarter_id=check.quarter.id, period_month=check.period, file_name=file_name[:255],
        file_sha256=check.sha256, file_size=check.size, file_format=check.file_format,
        mapping=check.mapping, row_count=check.row_count, rows_created=check.count("new"),
        rows_replaced=check.count("replace"), rows_unchanged=check.count("unchanged"),
        created_by_user_id=principal.user_id)
    db.add(record)
    db.flush()
    existing = {
        (v.sheet_id, v.kpi_definition_id): v
        for v in db.scalars(select(KpiMonthlyValue).where(
            KpiMonthlyValue.sheet_id.in_(sorted({line.sheet_id for line in check.lines})),
            KpiMonthlyValue.period_month == check.period)).all()
    }
    replaced = []
    for line in check.lines:
        if line.action == "unchanged":
            continue
        value = existing.get((line.sheet_id, line.kpi_definition_id))
        if value is None:
            value = KpiMonthlyValue(sheet_id=line.sheet_id, kpi_definition_id=line.kpi_definition_id,
                                    period_month=check.period)
            db.add(value)
        else:
            replaced.append({"sheet_id": line.sheet_id, "kpi": line.kpi_code,
                             "employee": line.employee_code,
                             "from": value.actual_value, "to": line.actual,
                             "inputs_from": value.input_values, "inputs_to": line.inputs})
        value.actual_value = line.actual
        value.input_values = line.inputs
        value.data_import_id = record.id
        value.source_row = line.row
        value.recorded_by_user_id = principal.user_id
    db.flush()
    record_event(
        db, actor=actor, action="kpi_data.import", entity_type="kpi_data_import",
        entity_id=record.id,
        summary=(f"Imported KPI data for {check.period:%B %Y} ({check.quarter.name}) from "
                 f"{file_name[:100]}: {record.rows_created} new, {record.rows_replaced} replaced, "
                 f"{record.rows_unchanged} unchanged"),
        changes={"file": file_name[:255], "sha256": check.sha256, "size": check.size,
                 "format": check.file_format, "period": check.period, "quarter_id": check.quarter.id,
                 "mapping": check.mapping, "rows": check.row_count,
                 "created": record.rows_created, "replaced": record.rows_replaced,
                 "unchanged": record.rows_unchanged, "replaced_values": replaced[:50]})
    return record


def list_imports(db: Session, principal: Principal, quarter_id: int) -> list[tuple[KpiDataImport, str | None]]:
    """Imports of a quarter (all of them for KPI writers; file metadata only)."""
    if principal.scope_for(Permission.KPI_WRITE).is_empty and principal.org_scope.is_empty:
        raise PermissionDeniedError("You cannot view KPI data imports.")
    rows = db.execute(
        select(KpiDataImport, User.username)
        .join(User, User.id == KpiDataImport.created_by_user_id, isouter=True)
        .where(KpiDataImport.cycle_quarter_id == quarter_id)
        .order_by(KpiDataImport.id.desc())).all()
    return [(r[0], r[1]) for r in rows]


def sheet_values(db: Session, sheet_id: int) -> list[KpiMonthlyValue]:
    return list(db.scalars(
        select(KpiMonthlyValue).where(KpiMonthlyValue.sheet_id == sheet_id)
        .options(joinedload(KpiMonthlyValue.kpi_definition), joinedload(KpiMonthlyValue.data_import))
        .order_by(KpiMonthlyValue.period_month, KpiMonthlyValue.kpi_definition_id)).all())


# --- template --------------------------------------------------------------------------------

_BOLD = Font(bold=True)
_FILL = PatternFill("solid", fgColor="DBE7F5")


def build_template(db: Session, principal: Principal, quarter_id: int, period: str) -> bytes:
    """A workbook pre-filled with one row per KPI on the published sheets in
    the caller's scope for the month; only the values need filling in."""
    quarter = get_quarter(db, quarter_id)
    month = parse_period(period)
    sheets = [s for s in sheets_in_scope(db, principal, quarter) if _active_in(s, month)]
    contexts = [_item_context(s, i) for s in sheets for i in s.items]
    extras = extra_fields(contexts)
    headers = [c.label for c in CONTRACT] + [c.label for c in extras]
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for col, heading in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=heading)
        cell.font, cell.fill = _BOLD, _FILL
        ws.column_dimensions[cell.column_letter].width = max(14, len(heading) + 4)
    ws.freeze_panes = "A2"
    row = 2
    for ctx in sorted(contexts, key=lambda c: (c.sheet.assignment.employee.employee_code,
                                               c.item.sort_order)):
        unit = ctx.sheet.assignment.plant_department
        values = [ctx.sheet.assignment.employee.employee_code, unit.plant.code, unit.department.code,
                  ctx.item.kpi_definition.code, f"{month:%Y-%m}", None, ctx.item.unit]
        for col, value in enumerate(values, 1):
            ws.cell(row=row, column=col, value=value)
        row += 1

    guide = wb.create_sheet("Instructions")
    lines = [f"Monthly KPI data — {quarter.name}, {month:%B %Y}", "",
             "One row per employee and KPI. Fill in Actual (and any extra input columns) and "
             "upload the file in the portal (Monthly KPI data). Nothing is saved until you "
             "confirm the preview.", "", "Columns:"]
    lines += [f"{c.label}{' *' if c.required else ''}: {c.description}" for c in (*CONTRACT, *extras)]
    lines += ["", "CSV files are accepted too (UTF-8, with the same column headings)."]
    for i, text in enumerate(lines, 1):
        cell = guide.cell(row=i, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True)
        if i == 1:
            cell.font = _BOLD
    guide.column_dimensions["A"].width = 110

    ref = wb.create_sheet("Reference")
    ref_headers = ["Employee Code", "Employee", "Assignment", "KPI", "KPI name", "Unit", "Target",
                   "Monthly values combine by", "Scoring matrix", "Needs"]
    for col, heading in enumerate(ref_headers, 1):
        ref.cell(row=1, column=col, value=heading).font = _BOLD
    for r, ctx in enumerate(contexts, 2):
        unit = ctx.sheet.assignment.plant_department
        measurement = ctx.item.measurement_snapshot or {}
        version = ctx.item.matrix_version
        needs = (["Actual"] if ctx.needs_actual else []) + [f.label for f in ctx.imported]
        for col, value in enumerate([
            ctx.sheet.assignment.employee.employee_code, ctx.sheet.assignment.employee.full_name,
            f"{unit.department.name} — {unit.plant.name}", ctx.item.kpi_definition.code,
            ctx.item.kpi_definition.name, ctx.item.unit,
            float(ctx.item.target_value) if ctx.item.target_value is not None else None,
            measurement.get("aggregation") or "—",
            f"{version.matrix.code} v{version.version_no}" if version else "Stage J rule",
            ", ".join(needs) or "—"], 1):
            ref.cell(row=r, column=col, value=value)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
