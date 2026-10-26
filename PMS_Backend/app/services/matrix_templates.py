"""Excel templates for matrix configurations: generate and parse.

A template is generated from metadata only: the matrix type's standard
fields (app/services/matrix_types.py) and, for an existing matrix, the field
definitions stored in its version. Adding a field to a type therefore adds
its columns here, to upload parsing, to the portal's forms and to previews,
with no other list to update.

Layout (TEMPLATE_VERSION 1):
  Instructions  how to fill in the workbook (generated)
  Matrix        settings: code, name, attributes, effective dates, the
                "when no rule matches" choice and default outputs
  Applies to    plant code / department code pairs ("ALL" = every one)
  Fields        the inputs and outputs, with a Use (Yes/No) column
  Rules         one row per rule; condition columns per input, value columns
                per output
  Test cases    representative inputs and the outputs they must give
  Codes         plant and department codes for reference
  _pms          very hidden: template kind, matrix type, schema and template
                versions, matrix code and base version

Row 1 of Fields, Rules and Test cases holds machine keys (hidden); row 2 the
headings people read. Parsing reads the keys, so editing headings is harmless.

Safety: files are parsed with openpyxl (through defusedxml), size- and
row-limited, and any cell containing a formula is refused: values are typed
in, never computed by the spreadsheet.
"""
from __future__ import annotations

import hashlib
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from app.models import Department, Plant
from app.services.errors import Issue, ValidationFailedError
from app.services.matrix_rules import _blank, decimal_text, jsonable, parse_bool, parse_decimal
from app.services.matrix_types import (
    DATA_TYPE_LABELS,
    RANGE_TYPES,
    DataType,
    MatrixTypeSpec,
)

TEMPLATE_VERSION = 1
TEMPLATE_KIND = "matrix_configuration"
MAX_BYTES = 1_000_000
MAX_ROWS = 1000
META_SHEET = "_pms"
FIELD_COLUMNS = ("use", "key", "label", "role", "data_type", "unit", "required", "min", "max",
                 "allowed_values", "source", "aggregation", "round_to", "description")
FIELD_HEADINGS = {
    "use": "Use (Yes/No)", "key": "Key *", "label": "Label *", "role": "Role (input/output) *",
    "data_type": "Data type *", "unit": "Unit", "required": "Required (Yes/No)",
    "min": "Minimum", "max": "Maximum", "allowed_values": "Allowed values (comma-separated)",
    "source": "Source", "aggregation": "Monthly aggregation", "round_to": "Round to decimals",
    "description": "Description",
}
LOWER_WORDS = {"more than (>)": False, "at least (≥)": True}
UPPER_WORDS = {"less than (<)": False, "at most (≤)": True}
NO_MATCH_WORDS = {"Use the default output": "default", "Stop with an error": "error"}
ALL = "ALL"

_BOLD = Font(bold=True)
_HEAD_FILL = PatternFill("solid", fgColor="DBE7F5")
_KEY_FONT = Font(color="808080", size=8)
_WRAP = Alignment(wrap_text=True, vertical="top")


# --- generation ------------------------------------------------------------------------------

@dataclass(frozen=True)
class TemplateContext:
    """What a template is generated for: a type, optionally a matrix version."""

    spec: MatrixTypeSpec
    matrix_code: str | None = None
    version_no: int | None = None
    name: str | None = None
    description: str | None = None
    attributes: Mapping[str, Any] = field(default_factory=dict)
    definition: Mapping[str, Any] | None = None
    effective_from: date | None = None
    effective_to: date | None = None
    applicability: Sequence[tuple[str, str]] = ()  # (plant code | ALL, department code | ALL)
    change_summary: str | None = None


def _field_rows(ctx: TemplateContext) -> list[dict[str, Any]]:
    """Fields to list: the version's own, then unused standard inputs."""
    rows: list[dict[str, Any]] = []
    stored = [f for f in (ctx.definition or {}).get("fields") or [] if isinstance(f, Mapping)]
    for f in stored:
        rows.append({**f, "use": True})
    listed = {r.get("key") for r in rows}
    for i, standard in enumerate(ctx.spec.standard_inputs):
        if standard.key not in listed:
            # A new matrix starts with the type's first standard input in use.
            rows.append({**standard.as_definition(), "use": not stored and i == 0})
    for output in ctx.spec.outputs:
        if output.key not in listed and output.key not in {r.get("key") for r in rows}:
            rows.append({**output.as_definition(), "use": True})
    inputs = [r for r in rows if r.get("role") == "input"]
    outputs = [r for r in rows if r.get("role") == "output"]
    return inputs + outputs


def _cell_text(value: Any) -> Any:
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    return value


def _header(ws: Worksheet, keys: Sequence[str], headings: Sequence[str],
            notes: Sequence[str | None] = ()) -> None:
    for col, (key, heading) in enumerate(zip(keys, headings), 1):
        key_cell = ws.cell(row=1, column=col, value=key)
        key_cell.font = _KEY_FONT
        cell = ws.cell(row=2, column=col, value=heading)
        cell.font, cell.fill, cell.alignment = _BOLD, _HEAD_FILL, _WRAP
        note = notes[col - 1] if col - 1 < len(notes) else None
        if note:
            cell.comment = Comment(note, "PMS")
        ws.column_dimensions[get_column_letter(col)].width = max(14, min(36, len(heading) + 4))
    ws.row_dimensions[1].hidden = True
    ws.freeze_panes = "A3"


def _list_validation(ws: Worksheet, values: Sequence[str], column: int, rows: int = 300) -> None:
    formula = '"' + ",".join(values) + '"'
    if len(formula) > 255:  # Excel's limit for an inline list
        return
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    letter = get_column_letter(column)
    dv.add(f"{letter}3:{letter}{rows + 2}")
    ws.add_data_validation(dv)


def _field_note(f: Mapping[str, Any]) -> str:
    parts = [f"{f.get('label')} — {DATA_TYPE_LABELS.get(DataType(f['data_type']), f['data_type'])}"]
    if f.get("unit"):
        parts.append(f"Unit: {f['unit']}")
    if f.get("min") is not None or f.get("max") is not None:
        parts.append(f"Accepted: {f.get('min') if f.get('min') is not None else '…'} to "
                     f"{f.get('max') if f.get('max') is not None else '…'}")
    if f.get("allowed_values"):
        parts.append(f"Values: {', '.join(f['allowed_values'])}")
    if f.get("description"):
        parts.append(str(f["description"]))
    return "\n".join(parts)


def rule_columns(fields: Sequence[Mapping[str, Any]]) -> list[tuple[str, str, str | None]]:
    """(key, heading, note) for each Rules column, from field metadata."""
    columns: list[tuple[str, str, str | None]] = [
        ("rule:priority", "Priority *", "A unique whole number; rules are listed by priority."),
        ("rule:label", "Rule label", None),
    ]
    for f in fields:
        if f.get("role") != "input":
            continue
        key, label, note = f["key"], f.get("label") or f["key"], _field_note(f)
        data_type = DataType(f["data_type"])
        if data_type in RANGE_TYPES:
            columns += [
                (f"cond:{key}:min", f"{label} from", note),
                (f"cond:{key}:min_inclusive", f"{label} from is", "more than (>) excludes the "
                 "value; at least (≥) includes it."),
                (f"cond:{key}:max", f"{label} to", note),
                (f"cond:{key}:max_inclusive", f"{label} to is", "less than (<) excludes the "
                 "value; at most (≤) includes it."),
            ]
        elif data_type == DataType.CATEGORY:
            columns.append((f"cond:{key}:values", f"{label} is one of", note))
        else:
            columns.append((f"cond:{key}:value", f"{label} is (Yes/No)", note))
    for f in fields:
        if f.get("role") == "output":
            columns.append((f"out:{f['key']}", f"{f.get('label') or f['key']} *", _field_note(f)))
    return columns


def case_columns(fields: Sequence[Mapping[str, Any]]) -> list[tuple[str, str, str | None]]:
    columns: list[tuple[str, str, str | None]] = [("case:label", "Test case", None)]
    for f in fields:
        if f.get("role") == "input":
            columns.append((f"in:{f['key']}", f.get("label") or f["key"], _field_note(f)))
    for f in fields:
        if f.get("role") == "output":
            columns.append((f"expect:{f['key']}", f"Expected {f.get('label') or f['key']} *",
                            _field_note(f)))
    return columns


def _rule_values(rule: Mapping[str, Any]) -> dict[str, Any]:
    values: dict[str, Any] = {"rule:priority": rule.get("priority"),
                              "rule:label": rule.get("label")}
    for key, condition in (rule.get("conditions") or {}).items():
        if not isinstance(condition, Mapping):
            continue
        if "values" in condition:
            values[f"cond:{key}:values"] = ", ".join(str(v) for v in condition["values"] or [])
        elif "value" in condition:
            values[f"cond:{key}:value"] = "Yes" if condition["value"] else "No"
        else:
            for side, words in (("min", LOWER_WORDS), ("max", UPPER_WORDS)):
                if condition.get(side) is None:
                    continue
                values[f"cond:{key}:{side}"] = _number_cell(condition[side])
                inclusive = condition.get(f"{side}_inclusive")
                if isinstance(inclusive, bool):
                    values[f"cond:{key}:{side}_inclusive"] = next(
                        w for w, inc in words.items() if inc is inclusive)
    for key, value in (rule.get("outputs") or {}).items():
        values[f"out:{key}"] = _number_cell(value)
    return values


def _exact_text(text: str) -> str | None:
    """Canonical decimal text, or None if `text` is not a plain number."""
    try:
        return decimal_text(parse_decimal(text))
    except ValueError:
        return None


def _number_cell(value: Any) -> Any:
    """Numbers become Excel numbers when that is exact; anything else stays text."""
    if not isinstance(value, str):
        return value
    canonical = _exact_text(value)
    if canonical is None:
        return value
    number = float(canonical)
    return number if _exact_text(repr(number)) == canonical else value


def build_template(ctx: TemplateContext, plants: Sequence[Plant],
                   departments: Sequence[Department]) -> bytes:
    spec = ctx.spec
    fields = _field_rows(ctx)
    wb = Workbook()

    guide = wb.active
    guide.title = "Instructions"
    lines = [
        (f"{spec.name} — configuration template", True),
        (spec.description, False),
        ("", False),
        ("How to fill in this workbook", True),
        ("1. Matrix: name the configuration (or keep the code of the matrix you are "
         "changing), and set the effective dates and what happens when no rule matches.", False),
        ("2. Applies to: one row per plant code and department code; ALL means every plant or "
         "every department (company-wide needs company-wide HR Head access).", False),
        ("3. Fields: set Use to Yes for each input the rules use. Outputs are required.", False),
        ("4. Rules: one row per rule. Leave a field's columns empty to accept any value. For "
         "numbers and dates give a lower and/or upper bound and say whether each bound is "
         "included: 'more than (>)' excludes the value itself, 'at least (≥)' includes it.", False),
        ("   Exactly one rule may match any input: rules must not overlap.", False),
        ("5. Test cases: representative inputs and the outputs they must give. They must pass "
         "before the configuration can be submitted to the CHRO.", False),
        ("6. Upload the workbook in the portal. It is checked first and nothing is saved until "
         "you confirm; then everything is saved as a draft, or nothing is.", False),
        ("", False),
        ("Enter plain values only: cells containing formulas are refused. Do not rename the "
         "sheets or edit the hidden first row of Fields, Rules and Test cases.", False),
    ]
    if spec.policy_note:
        lines += [("", False), (spec.policy_note, True)]
    for row, (text, bold) in enumerate(lines, 1):
        cell = guide.cell(row=row, column=1, value=text)
        cell.font, cell.alignment = (_BOLD if bold else Font()), _WRAP
    guide.column_dimensions["A"].width = 120

    settings = wb.create_sheet("Matrix")
    for col, heading in enumerate(("Key", "Setting", "Value", "Help"), 1):
        cell = settings.cell(row=1, column=col, value=heading)
        cell.font, cell.fill = _BOLD, _HEAD_FILL
    definition = ctx.definition or {}
    no_match = definition.get("no_match") if isinstance(definition.get("no_match"), Mapping) else None
    rows: list[tuple[str, str, Any, str]] = [
        ("matrix_code", "Matrix code", ctx.matrix_code, "Keep to change this matrix; leave "
         "empty to create a new one."),
        ("matrix_type", "Matrix type", spec.code, "Do not change."),
        ("name", "Name *", ctx.name, "Unique within this matrix type."),
        ("description", "Description", ctx.description, ""),
    ]
    for attribute in spec.attributes:
        rows.append((f"attr:{attribute.key}", f"{attribute.label} *",
                     ctx.attributes.get(attribute.key), attribute.description))
    rows += [
        ("effective_from", "Effective from *", ctx.effective_from, "First day the version "
         "applies (YYYY-MM-DD)."),
        ("effective_to", "Effective to", ctx.effective_to, "Last day; leave empty for "
         "open-ended."),
        ("no_match", "When no rule matches *",
         next((w for w, m in NO_MATCH_WORDS.items() if no_match and no_match.get("mode") == m),
              None), "Use the default output, or stop with an error."),
    ]
    for f in fields:
        if f.get("role") == "output":
            default = (no_match or {}).get("outputs") or {}
            rows.append((f"default:{f['key']}", f"Default {f.get('label')}",
                         _number_cell(default.get(f["key"])) if default.get(f["key"]) is not None
                         else None, "Used when no rule matches (with 'Use the default output')."))
    rows.append(("change_summary", "Change summary", ctx.change_summary,
                 "What changed in this version."))
    for row, (key, label, value, help_text) in enumerate(rows, 2):
        settings.cell(row=row, column=1, value=key).font = _KEY_FONT
        settings.cell(row=row, column=2, value=label).font = _BOLD
        cell = settings.cell(row=row, column=3, value=value)
        if isinstance(value, date):
            cell.number_format = "yyyy-mm-dd"
        settings.cell(row=row, column=4, value=help_text).alignment = _WRAP
        if key.startswith("attr:"):
            attribute = next(a for a in spec.attributes if f"attr:{a.key}" == key)
            dv = DataValidation(type="list", formula1='"' + ",".join(attribute.allowed_values) + '"',
                                allow_blank=False)
            dv.add(f"C{row}")
            settings.add_data_validation(dv)
        if key == "no_match":
            dv = DataValidation(type="list", formula1='"' + ",".join(NO_MATCH_WORDS) + '"')
            dv.add(f"C{row}")
            settings.add_data_validation(dv)
    settings.column_dimensions["A"].hidden = True
    for letter, width in zip("BCD", (28, 34, 70)):
        settings.column_dimensions[letter].width = width

    scope_sheet = wb.create_sheet("Applies to")
    for col, heading in enumerate(("Plant code", "Department code"), 1):
        cell = scope_sheet.cell(row=1, column=col, value=heading)
        cell.font, cell.fill = _BOLD, _HEAD_FILL
        scope_sheet.column_dimensions[get_column_letter(col)].width = 22
    for row, (plant, department) in enumerate(ctx.applicability, 2):
        scope_sheet.cell(row=row, column=1, value=plant)
        scope_sheet.cell(row=row, column=2, value=department)

    field_sheet = wb.create_sheet("Fields")
    _header(field_sheet, FIELD_COLUMNS, [FIELD_HEADINGS[c] for c in FIELD_COLUMNS])
    for row, f in enumerate(fields, 3):
        for col, key in enumerate(FIELD_COLUMNS, 1):
            value = f.get(key)
            if key == "required" and value is None:
                value = True
            field_sheet.cell(row=row, column=col, value=_cell_text(value))
    _list_validation(field_sheet, ["Yes", "No"], 1)
    _list_validation(field_sheet, ["input", "output"], 4)
    _list_validation(field_sheet, [t.value for t in DataType], 5)
    _list_validation(field_sheet, ["Yes", "No"], 7)
    field_sheet.column_dimensions["N"].width = 60

    rule_sheet = wb.create_sheet("Rules")
    columns = rule_columns(fields)
    _header(rule_sheet, [c[0] for c in columns], [c[1] for c in columns], [c[2] for c in columns])
    for row, rule in enumerate(definition.get("rules") or [], 3):
        if not isinstance(rule, Mapping):
            continue
        values = _rule_values(rule)
        for col, (key, _heading, _note) in enumerate(columns, 1):
            if key in values and values[key] is not None:
                rule_sheet.cell(row=row, column=col, value=values[key])
    for col, (key, _heading, _note) in enumerate(columns, 1):
        if key.endswith(":min_inclusive"):
            _list_validation(rule_sheet, list(LOWER_WORDS), col)
        elif key.endswith(":max_inclusive"):
            _list_validation(rule_sheet, list(UPPER_WORDS), col)
        elif key.endswith(":value"):
            _list_validation(rule_sheet, ["Yes", "No"], col)

    case_sheet = wb.create_sheet("Test cases")
    columns = case_columns(fields)
    _header(case_sheet, [c[0] for c in columns], [c[1] for c in columns], [c[2] for c in columns])
    for row, case in enumerate(definition.get("test_cases") or [], 3):
        if not isinstance(case, Mapping):
            continue
        values = {"case:label": case.get("label")}
        values.update({f"in:{k}": _cell_text(_number_cell(v))
                       for k, v in (case.get("inputs") or {}).items()})
        values.update({f"expect:{k}": _cell_text(_number_cell(v))
                       for k, v in (case.get("expected") or {}).items()})
        for col, (key, _heading, _note) in enumerate(columns, 1):
            if values.get(key) is not None:
                case_sheet.cell(row=row, column=col, value=values[key])

    codes = wb.create_sheet("Codes")
    codes.cell(row=1, column=1, value="Plant code").font = _BOLD
    codes.cell(row=1, column=2, value="Plant").font = _BOLD
    codes.cell(row=1, column=4, value="Department code").font = _BOLD
    codes.cell(row=1, column=5, value="Department").font = _BOLD
    for row, plant in enumerate(plants, 2):
        codes.cell(row=row, column=1, value=plant.code)
        codes.cell(row=row, column=2, value=plant.name)
    for row, department in enumerate(departments, 2):
        codes.cell(row=row, column=4, value=department.code)
        codes.cell(row=row, column=5, value=department.name)
    for letter, width in zip("ABCDE", (14, 28, 4, 18, 28)):
        codes.column_dimensions[letter].width = width

    meta = wb.create_sheet(META_SHEET)
    for row, (key, value) in enumerate((
        ("template_kind", TEMPLATE_KIND),
        ("template_version", TEMPLATE_VERSION),
        ("matrix_type", spec.code),
        ("schema_version", spec.schema_version),
        ("matrix_code", ctx.matrix_code),
        ("base_version_no", ctx.version_no),
        ("generated_at", datetime.now().replace(microsecond=0).isoformat()),
    ), 1):
        meta.cell(row=row, column=1, value=key)
        meta.cell(row=row, column=2, value=value)
    meta.sheet_state = "veryHidden"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- parsing ---------------------------------------------------------------------------------

CONDITION_PARTS = frozenset({"min", "min_inclusive", "max", "max_inclusive", "values", "value"})


@dataclass
class Located:
    """A problem with its place in the workbook, e.g. "Rules sheet, row 5"."""

    location: str
    message: str


@dataclass
class ParsedWorkbook:
    meta: dict[str, Any]
    settings: dict[str, Any]  # setting key -> raw cell value
    applicability: list[tuple[str, str, int]]  # (plant code | ALL, department code | ALL, row)
    definition: dict[str, Any]
    field_rows: dict[int, int]  # index in definition["fields"] -> sheet row
    rule_rows: dict[int, int]
    case_rows: dict[int, int]
    problems: list[Located]
    sha256: str


def _unreadable(message: str) -> ValidationFailedError:
    return ValidationFailedError([Issue(message)])


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip() or None


def _scalar(value: Any) -> Any:
    """A cell value as exact input for the rule engine (numbers as text)."""
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return _exact_text(repr(value)) or repr(value)
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _int_or_text(value: Any) -> Any:
    if isinstance(value, bool) or value is None:
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = str(value).strip()
    return int(text) if text.isdigit() else text


def _data_type(value: Any) -> DataType | None:
    try:
        return DataType(value)
    except ValueError:
        return None


def _rows(ws: Worksheet, sheet: str, problems: list[Located]) -> list[tuple[int, list[Any]]]:
    """(row number, plain values). Formula cells are refused, never evaluated."""
    rows = []
    for number, row in enumerate(ws.iter_rows(min_row=1), 1):
        if number > MAX_ROWS + 2:
            raise _unreadable(f"The {sheet} sheet has more than {MAX_ROWS} rows.")
        values = []
        for col, cell in enumerate(row, 1):
            value = getattr(cell, "value", None)
            if getattr(cell, "data_type", None) == "f":
                problems.append(Located(f"{sheet} sheet, cell {get_column_letter(col)}{number}",
                                        "Contains a formula; enter a plain value."))
                value = None
            values.append(value.strip() if isinstance(value, str) else value)
        rows.append((number, values))
    return rows


def _keyed(ws: Worksheet, sheet: str, problems: list[Located]
           ) -> tuple[list[str], list[tuple[int, dict[str, Any]]]]:
    """Machine keys from the hidden first row, and the non-empty data rows."""
    rows = _rows(ws, sheet, problems)
    keys = [(_text(v) or "") for v in rows[0][1]] if rows else []
    records = []
    for number, values in rows[2:]:  # row 1: keys, row 2: headings
        record = {k: v for k, v in zip(keys, values) if k and not _blank(v)}
        if record:
            records.append((number, record))
    return keys, records


def _parse_fields(records: list[tuple[int, dict[str, Any]]], problems: list[Located]):
    fields: list[dict[str, Any]] = []
    rows: dict[int, int] = {}
    listed: dict[str, str] = {}
    used: set[str] = set()
    for number, record in records:
        where = f"Fields sheet, row {number}"
        unknown = sorted(set(record) - set(FIELD_COLUMNS))
        if unknown:
            problems.append(Located(where, f"Unknown column(s): {', '.join(unknown)}."))
        key = _text(record.get("key"))
        if not key:
            problems.append(Located(where, "Give the field a key."))
            continue
        if key in listed:
            problems.append(Located(where, f"The field '{key}' is listed twice."))
            continue
        listed[key] = (_text(record.get("role")) or "").lower()
        try:
            in_use = parse_bool(record.get("use", "Yes"))
        except ValueError:
            problems.append(Located(where, "Use must be Yes or No."))
            continue
        if not in_use:
            continue
        required_raw = record.get("required")
        try:
            required = True if _blank(required_raw) else parse_bool(required_raw)
        except ValueError:
            problems.append(Located(where, "Required must be Yes or No."))
            required = True
        allowed = _text(record.get("allowed_values"))
        used.add(key)
        rows[len(fields)] = number
        fields.append({
            "key": key,
            "label": _text(record.get("label")),
            "role": listed[key],
            "data_type": (_text(record.get("data_type")) or "").upper(),
            "unit": _text(record.get("unit")),
            "description": _text(record.get("description")),
            "required": required,
            "min": _scalar(record.get("min")),
            "max": _scalar(record.get("max")),
            "allowed_values": [v.strip() for v in allowed.split(",") if v.strip()] if allowed else None,
            "source": (_text(record.get("source")) or "").upper() or None,
            "aggregation": (_text(record.get("aggregation")) or "").upper() or None,
            "round_to": _int_or_text(record.get("round_to")),
        })
    return fields, rows, listed, used


def _check_rule_columns(keys: list[str], listed: dict[str, str], used: set[str],
                        problems: list[Located]) -> None:
    for col, key in enumerate(keys, 1):
        if not key:
            continue
        parts = key.split(":")
        known = (
            (parts[0] == "rule" and len(parts) == 2 and parts[1] in ("priority", "label"))
            or (parts[0] == "out" and len(parts) == 2 and listed.get(parts[1]) == "output")
            or (parts[0] == "cond" and len(parts) == 3 and listed.get(parts[1]) == "input"
                and parts[2] in CONDITION_PARTS)
        )
        if not known:
            problems.append(Located(f"Rules sheet, column {get_column_letter(col)}",
                                    f"The column '{key}' does not match a field on the Fields "
                                    "sheet. Download the template for this matrix again after "
                                    "changing its fields."))
    for key in sorted(used):
        prefix = f"cond:{key}:" if listed.get(key) == "input" else f"out:{key}"
        if not any(k == prefix or k.startswith(prefix) for k in keys if k):
            problems.append(Located("Rules sheet", f"There are no columns for '{key}'. Download "
                                                   "the template for this matrix again after "
                                                   "adding fields."))


def _parse_rules(records: list[tuple[int, dict[str, Any]]], used: set[str],
                 types: dict[str, DataType | None], problems: list[Located]):
    rules: list[dict[str, Any]] = []
    rows: dict[int, int] = {}
    for number, record in records:
        where = f"Rules sheet, row {number}"
        rule: dict[str, Any] = {"priority": _int_or_text(record.get("rule:priority")),
                                "label": _text(record.get("rule:label")),
                                "conditions": {}, "outputs": {}}
        for key, value in record.items():
            parts = key.split(":")
            if parts[0] == "out" and len(parts) == 2:
                rule["outputs"][parts[1]] = _scalar(value)
                continue
            if parts[0] != "cond" or len(parts) != 3:
                continue
            field_key, part = parts[1], parts[2]
            if field_key not in used:
                problems.append(Located(where, f"Uses '{field_key}', which is not marked "
                                               "Use = Yes on the Fields sheet."))
                continue
            condition = rule["conditions"].setdefault(field_key, {})
            if part in ("min_inclusive", "max_inclusive"):
                words = LOWER_WORDS if part == "min_inclusive" else UPPER_WORDS
                flag = _inclusive_flag(str(value), words)
                if flag is None:
                    problems.append(Located(where, f"'{value}' is not one of: "
                                                   f"{', '.join(words)}."))
                condition[part] = flag
            elif part == "values":
                condition["values"] = [v.strip() for v in str(value).split(",") if v.strip()]
            elif part == "value":
                try:
                    condition["value"] = parse_bool(value)
                except ValueError:
                    problems.append(Located(where, f"'{value}' must be Yes or No."))
            elif part in ("min", "max"):
                condition[part] = _scalar(value)
        for key, condition in rule["conditions"].items():
            if types.get(key) in RANGE_TYPES:
                for part in ("min", "max", "min_inclusive", "max_inclusive"):
                    condition.setdefault(part, None)
        rows[len(rules)] = number
        rules.append(rule)
    return rules, rows


def _parse_cases(records: list[tuple[int, dict[str, Any]]], used: set[str],
                 problems: list[Located]):
    cases: list[dict[str, Any]] = []
    rows: dict[int, int] = {}
    for number, record in records:
        inputs, expected = {}, {}
        for key, value in record.items():
            parts = key.split(":")
            if parts[0] == "in" and len(parts) == 2:
                if parts[1] in used:
                    inputs[parts[1]] = _scalar(value)
                else:
                    problems.append(Located(f"Test cases sheet, row {number}",
                                            f"'{parts[1]}' is not marked Use = Yes on the "
                                            "Fields sheet."))
            elif parts[0] == "expect" and len(parts) == 2:
                expected[parts[1]] = _scalar(value)
        rows[len(cases)] = number
        cases.append({"label": _text(record.get("case:label")) or f"Row {number}",
                      "inputs": inputs, "expected": expected})
    return cases, rows


def _parse_no_match(settings: dict[str, Any], problems: list[Located]) -> dict[str, Any] | None:
    word = _text(settings.get("no_match"))
    mode = next((m for w, m in NO_MATCH_WORDS.items() if word and w.lower() == word.lower()), None)
    if mode == "default":
        outputs = {key.split(":", 1)[1]: _scalar(value) for key, value in settings.items()
                   if key.startswith("default:") and not _blank(value)}
        return {"mode": "default", "outputs": outputs}
    if mode == "error":
        return {"mode": "error"}
    if word:
        problems.append(Located("Matrix sheet", f"'{word}' is not one of: "
                                                f"{', '.join(NO_MATCH_WORDS)}."))
    return None


def parse_workbook(data: bytes) -> ParsedWorkbook:
    """Read a filled-in template. Unreadable files raise; content problems are
    returned (with locations) for the validation report."""
    if not data:
        raise _unreadable("The file is empty.")
    if len(data) > MAX_BYTES:
        raise _unreadable("The file is larger than 1 MB.")
    if not data.startswith(b"PK\x03\x04"):
        raise _unreadable("Upload an Excel workbook (.xlsx) generated by the portal.")
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=False)
    except (InvalidFileException, BadZipFile, KeyError, OSError, ValueError):
        raise _unreadable("The file could not be read as an Excel workbook (.xlsx).") from None
    problems: list[Located] = []
    try:
        if META_SHEET not in wb.sheetnames:
            raise _unreadable("This workbook is not a PMS matrix template (its hidden version "
                              "sheet is missing). Download a fresh template from the portal.")
        meta: dict[str, Any] = {}
        for _number, values in _rows(wb[META_SHEET], META_SHEET, problems):
            key = _text(values[0]) if values else None
            if key:
                meta[key] = values[1] if len(values) > 1 else None
        missing = [s for s in ("Matrix", "Applies to", "Fields", "Rules", "Test cases")
                   if s not in wb.sheetnames]
        if missing:
            raise _unreadable(f"The workbook is missing sheet(s): {', '.join(missing)}. Download "
                              "a fresh template.")
        settings: dict[str, Any] = {}
        for _number, values in _rows(wb["Matrix"], "Matrix", problems)[1:]:
            key = _text(values[0]) if values else None
            if key:
                settings[key] = values[2] if len(values) > 2 else None
        applicability = []
        for number, values in _rows(wb["Applies to"], "Applies to", problems)[1:]:
            plant = _text(values[0]) if values else None
            department = _text(values[1]) if len(values) > 1 else None
            if plant or department:
                applicability.append(((plant or "").upper(), (department or "").upper(), number))
        _field_keys, field_records = _keyed(wb["Fields"], "Fields", problems)
        rule_keys, rule_records = _keyed(wb["Rules"], "Rules", problems)
        _case_keys, case_records = _keyed(wb["Test cases"], "Test cases", problems)
    finally:
        wb.close()

    fields, field_rows, listed, used = _parse_fields(field_records, problems)
    types = {f["key"]: _data_type(f["data_type"]) for f in fields}
    _check_rule_columns(rule_keys, listed, used, problems)
    rules, rule_rows = _parse_rules(rule_records, used, types, problems)
    cases, case_rows = _parse_cases(case_records, used, problems)
    definition = {"fields": fields, "rules": rules, "no_match": _parse_no_match(settings, problems),
                  "test_cases": cases}
    return ParsedWorkbook(meta=meta, settings=settings, applicability=applicability,
                          definition=jsonable(definition), field_rows=field_rows,
                          rule_rows=rule_rows, case_rows=case_rows, problems=problems,
                          sha256=hashlib.sha256(data).hexdigest())


def _inclusive_flag(text: str, words: Mapping[str, bool]) -> bool | None:
    """Whether a bound is included, from the template's words (or common
    shorthand). A lower-bound word in an upper-bound column is refused."""
    cleaned = text.strip().lower()
    for word, flag in words.items():
        if cleaned == word.lower():
            return flag
    lower = {">": False, ">=": True, "≥": True, "more than": False, "at least": True}
    upper = {"<": False, "<=": True, "≤": True, "less than": False, "at most": True}
    neutral = {"exclusive": False, "inclusive": True, "excluded": False, "included": True}
    is_lower_column = set(words) == set(LOWER_WORDS)
    if cleaned in neutral:
        return neutral[cleaned]
    return (lower if is_lower_column else upper).get(cleaned)
